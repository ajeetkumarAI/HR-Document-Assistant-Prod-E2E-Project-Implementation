"""Document loaders: PDF, DOCX, TXT, Markdown, HTML, CSV -> ``Document`` objects with rich metadata.

Headings are converted to Markdown (``#``) wherever the format exposes them, so the
structure-aware chunker can keep sections together and record the section path.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import Any

from src.ingestion.metadata import (
    file_metadata,
    load_sidecar,
    normalize_metadata,
    split_front_matter,
)
from src.ingestion.models import Document
from src.utils.exceptions import IngestionError, UnsupportedFileTypeError
from src.utils.helpers import normalize_text
from src.utils.logger import get_logger

logger = get_logger(__name__)

# Each parser returns: list of (text, extra_metadata)
Parser = Callable[[bytes], list[tuple[str, dict[str, Any]]]]


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _parse_text(content: bytes) -> list[tuple[str, dict[str, Any]]]:
    meta, body = split_front_matter(_decode(content))
    return [(body, meta)]


def _parse_pdf(content: bytes) -> list[tuple[str, dict[str, Any]]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    info = reader.metadata or {}
    doc_meta = {k: v for k, v in {"title": info.get("/Title"), "author": info.get("/Author")}.items() if v}
    pages: list[tuple[str, dict[str, Any]]] = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append((text, {**doc_meta, "page": i, "total_pages": len(reader.pages)}))
    if not pages:
        raise IngestionError("PDF has no extractable text (scanned?). Run OCR first, e.g. `ocrmypdf in.pdf out.pdf`.")
    return pages


def _parse_docx(content: bytes) -> list[tuple[str, dict[str, Any]]]:
    import docx

    d = docx.Document(io.BytesIO(content))
    lines: list[str] = []
    for para in d.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append(f"{'#' * min(int(level), 6)} {text}")
        elif style == "title":
            lines.append(f"# {text}")
        else:
            lines.append(text)
    for table in d.tables:  # tables -> pipe rows so the values stay together
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    props = d.core_properties
    meta = {k: v for k, v in {"title": props.title, "author": props.author}.items() if v}
    return [("\n\n".join(lines), meta)]


def _parse_html(content: bytes) -> list[tuple[str, dict[str, Any]]]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(content), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    for level in range(1, 7):
        for h in soup.find_all(f"h{level}"):
            h.replace_with(f"\n{'#' * level} {h.get_text(strip=True)}\n")
    title = soup.title.get_text(strip=True) if soup.title else None
    return [(soup.get_text("\n"), {"title": title} if title else {})]


def _parse_csv(content: bytes) -> list[tuple[str, dict[str, Any]]]:
    """Each row becomes 'column: value' lines - keeps tabular facts (e.g. holiday lists) retrievable."""
    reader = csv.DictReader(io.StringIO(_decode(content)))
    rows = ["\n".join(f"{k}: {v}" for k, v in row.items() if k and v not in (None, "")) for row in reader]
    return [("\n\n".join(rows), {"rows": len(rows)})]


PARSERS: dict[str, Parser] = {
    ".txt": _parse_text,
    ".md": _parse_text,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".html": _parse_html,
    ".htm": _parse_html,
    ".csv": _parse_csv,
}


class DocumentLoader:
    def __init__(self, default_metadata: dict[str, Any] | None = None, supported: list[str] | None = None):
        self.default_metadata = default_metadata or {}
        self.supported = {e.lower() for e in (supported or PARSERS.keys())}

    def load_bytes(
        self,
        content: bytes,
        filename: str,
        extra_metadata: dict[str, Any] | None = None,
        path: Path | None = None,
    ) -> list[Document]:
        ext = Path(filename).suffix.lower()
        if ext not in self.supported or ext not in PARSERS:
            raise UnsupportedFileTypeError(f"Unsupported file type '{ext}' for {filename}")
        try:
            parts = PARSERS[ext](content)
        except IngestionError:
            raise
        except Exception as exc:
            raise IngestionError(f"Failed to parse {filename}: {exc}") from exc

        base = {
            **self.default_metadata,
            **file_metadata(path or Path(filename), content, logical_name=filename),
            **(load_sidecar(path) if path else {}),
        }
        docs: list[Document] = []
        for text, part_meta in parts:
            text = normalize_text(text)
            if not text:
                continue
            meta = normalize_metadata({**base, **part_meta, **(extra_metadata or {})})
            docs.append(Document(text=text, metadata=meta))
        if not docs:
            raise IngestionError(f"No text content found in {filename}")
        logger.info("Loaded document", extra={"source": filename, "parts": len(docs)})
        return docs

    def load_file(self, path: str | Path, extra_metadata: dict[str, Any] | None = None) -> list[Document]:
        path = Path(path)
        return self.load_bytes(path.read_bytes(), path.name, extra_metadata, path=path)

    def iter_directory(self, directory: str | Path) -> Iterator[Path]:
        directory = Path(directory)
        for path in sorted(directory.rglob("*")):
            if path.is_file() and path.suffix.lower() in self.supported and not path.name.endswith(".meta.yaml"):
                yield path
