from __future__ import annotations

import io

import pytest

from src.chunking.chunker import Chunker
from src.ingestion.loader import DocumentLoader
from src.ingestion.metadata import normalize_metadata
from src.ingestion.models import Document
from src.utils.exceptions import UnsupportedFileTypeError
from src.utils.helpers import count_tokens
from tests.conftest import SAMPLE_DIR


def test_markdown_front_matter_becomes_metadata() -> None:
    docs = DocumentLoader().load_file(SAMPLE_DIR / "leave_policy.md")
    meta = docs[0].metadata
    assert meta["category"] == "leave"
    assert meta["access_level"] == "public"
    assert meta["effective_ts"] == 20260101
    assert "sick leave" in meta["tags"]
    assert not docs[0].text.startswith("---")  # front matter stripped from body


def test_csv_uses_sidecar_metadata() -> None:
    docs = DocumentLoader().load_file(SAMPLE_DIR / "holidays_2026.csv")
    assert docs[0].metadata["category"] == "holidays"
    assert "holiday: Diwali" in docs[0].text


def test_docx_headings_preserved() -> None:
    import docx

    d = docx.Document()
    d.add_heading("Gratuity Policy", level=1)
    d.add_heading("Eligibility", level=2)
    d.add_paragraph("Gratuity is payable after 5 years of continuous service.")
    buf = io.BytesIO()
    d.save(buf)
    docs = DocumentLoader().load_bytes(buf.getvalue(), "gratuity.docx", {"category": "benefits"})
    assert "## Eligibility" in docs[0].text
    assert docs[0].metadata["category"] == "benefits"


def test_unsupported_extension_rejected() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        DocumentLoader().load_bytes(b"x", "malware.exe")


def test_metadata_normalisation() -> None:
    meta = normalize_metadata({"access_level": "TOP-SECRET", "tags": "A, b ,a", "effective_date": "15/04/2026"})
    assert meta["access_level"] == "public"  # unknown levels downgrade safely
    assert meta["tags"] == ["a", "b"]
    assert meta["effective_ts"] == 20260415


def test_structure_aware_chunking_keeps_section_path() -> None:
    docs = DocumentLoader().load_file(SAMPLE_DIR / "leave_policy.md")
    chunks = Chunker(chunk_size=200, chunk_overlap=30, min_chunk_tokens=20).split_documents(docs)
    sections = {c.metadata["section"] for c in chunks}
    assert "Leave Policy > Sick Leave" in sections
    sick = next(c for c in chunks if c.metadata["section"] == "Leave Policy > Sick Leave")
    assert sick.embed_text.startswith("Document: Leave Policy\nSection: Leave Policy > Sick Leave")
    assert "medical certificate" in sick.text


def test_long_sections_split_within_budget_with_overlap() -> None:
    sentences = [f"Rule number {i} says employees must follow procedure {i} carefully." for i in range(120)]
    doc = Document(text="# Big\n\n" + " ".join(sentences), metadata={"doc_id": "d1", "title": "Big"})
    chunker = Chunker(chunk_size=120, chunk_overlap=30, min_chunk_tokens=10)
    chunks = chunker.split_document(doc)
    assert len(chunks) > 3
    assert all(count_tokens(c.text) <= 120 for c in chunks)
    # overlap: the last sentence of chunk i appears at the start of chunk i+1
    last_sentence = chunks[0].text.split("\n")[-1]
    assert chunks[1].text.startswith(last_sentence)


def test_chunk_ids_are_deterministic() -> None:
    docs = DocumentLoader().load_file(SAMPLE_DIR / "code_of_conduct.md")
    a = [c.chunk_id for c in Chunker().split_documents(docs)]
    b = [c.chunk_id for c in Chunker().split_documents(docs)]
    assert a == b and len(set(a)) == len(a)
